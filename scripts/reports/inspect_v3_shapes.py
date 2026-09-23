import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from pptx import Presentation

v3_path = Path("/Users/guobiao/PRO/me/yoloExample/docs/presentations/2026-09-23_YOLO门店陈列识别_COO_CTO汇报_v3.pptx")
prs = Presentation(v3_path)

print(f"Total slides in v3: {len(prs.slides)}")

slide2 = prs.slides[1]
print("\n--- Slide 2 Shapes in v3 ---")
for idx, shape in enumerate(slide2.shapes):
    text_preview = ""
    if shape.has_text_frame:
        text_preview = shape.text_frame.text.replace("\n", " ")[:40]
    print(f"Shape {idx}: '{shape.name}', Type={shape.shape_type}, Text='{text_preview}'")
    if hasattr(shape, "adjustments") and len(shape.adjustments) > 0:
        for adj_idx, adj_val in enumerate(shape.adjustments):
            print(f"   Adjustment {adj_idx} = {adj_val}")

# Also inspect slide 2 XML directly to get exact XML prstGeom and avLst
with zipfile.ZipFile(v3_path, 'r') as z:
    content = z.read('ppt/slides/slide2.xml')
    root = ET.fromstring(content)
    print("\n--- Slide 2 XML Shapes & Geometry Adjustments ---")
    for sp in root.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}sp'):
        name = sp.find('.//{*}cNvPr').attrib.get('name', 'Unknown')
        geom = sp.find('.//{*}prstGeom')
        geom_name = geom.attrib.get('prst') if geom is not None else 'None'
        avLst = geom.find('{*}avLst') if geom is not None else None
        gd_vals = []
        if avLst is not None:
            for gd in avLst.iter('{*}gd'):
                gd_vals.append((gd.attrib.get('name'), gd.attrib.get('fmla')))
        print(f"Shape Name: '{name}', Geom: '{geom_name}', AvLst: {gd_vals}")

