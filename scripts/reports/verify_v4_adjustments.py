from pptx import Presentation
from pathlib import Path

v4_path = Path("/Users/guobiao/PRO/me/yoloExample/docs/presentations/2026-09-23_YOLO门店陈列识别_COO_CTO汇报_v4.pptx")
prs = Presentation(v4_path)

print(f"Total slides in v4: {len(prs.slides)}")
for s_idx, slide in enumerate(prs.slides, start=1):
    rounded_shapes = []
    for sh_idx, shape in enumerate(slide.shapes):
        if hasattr(shape, "adjustments") and len(shape.adjustments) > 0:
            rounded_shapes.append((shape.name, round(shape.adjustments[0], 5)))
    print(f"Slide {s_idx}: {len(rounded_shapes)} rounded shapes -> adjustments: {set(v for _, v in rounded_shapes)}")

