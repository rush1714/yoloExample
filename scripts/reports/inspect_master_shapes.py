from pptx import Presentation
from pathlib import Path

ref_path = Path("/Users/guobiao/DOC/森大2.0/7.项目管理/CRM2.0开发资源及森大业务员访销方案_v0.6.pptx")
prs = Presentation(ref_path)

print(f"Number of slide masters: {len(prs.slide_masters)}")
for m_idx, master in enumerate(prs.slide_masters):
    print(f"Master {m_idx}: shapes count = {len(master.shapes)}")
    for s in master.shapes:
        print(f"  Master Shape: {s.name}, Type={s.shape_type}")

for l_idx, layout in enumerate(prs.slide_layouts):
    print(f"Layout {l_idx} ({layout.name}): shapes count = {len(layout.shapes)}")
    for s in layout.shapes:
        print(f"  Layout Shape: {s.name}, Type={s.shape_type}")
