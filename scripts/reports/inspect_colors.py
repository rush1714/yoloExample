from pptx import Presentation
from pathlib import Path

ref_path = Path("/Users/guobiao/DOC/森大2.0/7.项目管理/CRM2.0开发资源及森大业务员访销方案_v0.6.pptx")
prs = Presentation(ref_path)

slide1 = prs.slides[0]
for idx, shape in enumerate(slide1.shapes):
    line_col = shape.line.color.rgb if hasattr(shape, "line") and shape.line and shape.line.color and shape.line.color.type == 1 else "None"
    fill_col = "None"
    if hasattr(shape, "fill") and shape.fill and shape.fill.type == 1:
        fill_col = shape.fill.fore_color.rgb
    print(f"Shape {idx}: {shape.name}, Type={shape.shape_type}, Fill={fill_col}, Line={line_col}")
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            f_col = p.font.color.rgb if p.font and p.font.color and p.font.color.type == 1 else "Default"
            print(f"   Paragraph: '{p.text}' Font: {p.font.name} Size: {p.font.size} Color: {f_col}")

