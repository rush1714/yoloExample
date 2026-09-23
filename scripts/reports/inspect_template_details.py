from pptx import Presentation
from pathlib import Path

ref_path = Path("/Users/guobiao/DOC/森大2.0/7.项目管理/CRM2.0开发资源及森大业务员访销方案_v0.6.pptx")
prs = Presentation(ref_path)

print("Slide width:", prs.slide_width.inches, "height:", prs.slide_height.inches)
print("Number of slide layouts in master:", len(prs.slide_layouts))

for idx, layout in enumerate(prs.slide_layouts):
    print(f"Layout {idx}: {layout.name}")
    for s in layout.shapes:
        print(f"  Shape: {s.name}, Type={s.shape_type}")

print("\n--- Inspecting Slide 1 from original file ---")
s1 = prs.slides[0]
for s in s1.shapes:
    print(f"Shape: {s.name}, type: {s.shape_type}")
    if hasattr(s, "image"):
        print(f"  Image shape! {s.name}")

print("\n--- Inspecting Slide 2 from original file ---")
s2 = prs.slides[1]
for s in s2.shapes:
    print(f"Shape: {s.name}, type: {s.shape_type}")
    if hasattr(s, "image"):
        print(f"  Image shape! {s.name}")

