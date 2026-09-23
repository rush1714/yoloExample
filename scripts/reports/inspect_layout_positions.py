from pptx import Presentation
from pathlib import Path

ref_path = Path("/Users/guobiao/DOC/森大2.0/7.项目管理/CRM2.0开发资源及森大业务员访销方案_v0.6.pptx")
prs = Presentation(ref_path)

# Let's check master shape positions
master = prs.slide_masters[0]
print("--- Slide Master Shapes & Positions (Inches) ---")
for s in master.shapes:
    l = round(s.left.inches, 2)
    t = round(s.top.inches, 2)
    w = round(s.width.inches, 2)
    h = round(s.height.inches, 2)
    print(f"Master Shape: '{s.name}' | Left={l}, Top={t}, Width={w}, Height={h}, Type={s.shape_type}")

print("\n--- Layout 0 (Content) Placeholders ---")
l0 = prs.slide_layouts[0]
for s in l0.placeholders:
    l = round(s.left.inches, 2)
    t = round(s.top.inches, 2)
    w = round(s.width.inches, 2)
    h = round(s.height.inches, 2)
    print(f"Placeholder: '{s.name}' | Left={l}, Top={t}, Width={w}, Height={h}")

print("\n--- Layout 2 (Cover) Placeholders & Shapes ---")
l2 = prs.slide_layouts[2]
for s in l2.shapes:
    l = round(s.left.inches, 2)
    t = round(s.top.inches, 2)
    w = round(s.width.inches, 2)
    h = round(s.height.inches, 2)
    print(f"Cover Shape: '{s.name}' | Left={l}, Top={t}, Width={w}, Height={h}")

