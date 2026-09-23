from pptx import Presentation
from pathlib import Path

v3_path = Path("/Users/guobiao/PRO/me/yoloExample/docs/presentations/2026-09-23_YOLO门店陈列识别_COO_CTO汇报_v3.pptx")
prs = Presentation(v3_path)

for s_idx, slide in enumerate(prs.slides, start=1):
    print(f"\n--- Slide {s_idx} ---")
    for sh_idx, shape in enumerate(slide.shapes):
        if hasattr(shape, "adjustments") and len(shape.adjustments) > 0:
            print(f"  Shape {sh_idx}: '{shape.name}' Type={shape.shape_type} Adj0={shape.adjustments[0]}")
