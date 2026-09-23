import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ref_path = Path("/Users/guobiao/DOC/森大2.0/7.项目管理/CRM2.0开发资源及森大业务员访销方案_v0.6.pptx")
if not ref_path.is_file():
    print("File not found!")
    sys.exit(1)

prs = Presentation(ref_path)
print(f"Slide width: {prs.slide_width.inches} inches, height: {prs.slide_height.inches} inches")
print(f"Total slides: {len(prs.slides)}")

slide1 = prs.slides[0]
print("--- Slide 1 Shapes ---")
for shape in slide1.shapes:
    print(f"Shape name: {shape.name}, type: {shape.shape_type}")
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            text = p.text.strip()
            if text:
                font_name = p.font.name if p.font else "Unknown"
                font_size = p.font.size.pt if p.font and p.font.size else "Unknown"
                color = "Unknown"
                if p.font and p.font.color:
                    try:
                        color = str(p.font.color.rgb)
                    except Exception:
                        color = str(p.font.color.type)
                print(f"  Text: {text} | Font: {font_name} | Size: {font_size} | Color: {color}")
    if hasattr(shape, "fill") and shape.fill:
        try:
            print(f"  Fill type: {shape.fill.type}, ForeColor: {shape.fill.fore_color.rgb if shape.fill.type==1 else 'N/A'}")
        except Exception as e:
            pass

print("\n--- Slide 2 Shapes ---")
if len(prs.slides) > 1:
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    font_name = p.font.name if p.font else "Unknown"
                    font_size = p.font.size.pt if p.font and p.font.size else "Unknown"
                    color = "Unknown"
                    if p.font and p.font.color:
                        try:
                            color = str(p.font.color.rgb)
                        except Exception:
                            pass
                    print(f"  Text: {text[:50]} | Font: {font_name} | Size: {font_size} | Color: {color}")

