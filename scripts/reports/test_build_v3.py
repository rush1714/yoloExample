from pptx import Presentation
from pathlib import Path

ref_path = Path("/Users/guobiao/DOC/森大2.0/7.项目管理/CRM2.0开发资源及森大业务员访销方案_v0.6.pptx")
prs = Presentation(ref_path)

# Clear existing slides
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

# Add slide 1 with Layout 0 (to keep white background and master header/footer)
s1 = prs.slides.add_slide(prs.slide_layouts[0])
print("Added slide 1")

test_out = Path("/Users/guobiao/PRO/me/yoloExample/.tmp/test_v3.pptx")
prs.save(test_out)
print("Saved to", test_out)
