from pathlib import Path
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

v4_path = Path("/Users/guobiao/PRO/me/yoloExample/docs/presentations/2026-09-23_YOLO门店陈列识别_COO_CTO汇报_v4.pptx")
v5_path = Path("/Users/guobiao/PRO/me/yoloExample/docs/presentations/2026-09-23_YOLO门店陈列识别_COO_CTO汇报_v5.pptx")

prs = Presentation(v4_path)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

def set_font_on_rpr(rPr, font_name="Microsoft YaHei"):
    """在 DrawingML rPr 节点上设置 latin, ea, cs 等全系字体为指定字体。"""
    # 查找或创建 latin
    latin = rPr.find(f"{{{A_NS}}}latin")
    if latin is None:
        latin = parse_xml(f'<a:latin {nsdecls("a")} typeface="{font_name}"/>')
        rPr.append(latin)
    else:
        latin.set("typeface", font_name)

    # 查找或创建 ea (East Asian font, 确保中文字体正确渲染为微软雅黑)
    ea = rPr.find(f"{{{A_NS}}}ea")
    if ea is None:
        ea = parse_xml(f'<a:ea {nsdecls("a")} typeface="{font_name}"/>')
        rPr.append(ea)
    else:
        ea.set("typeface", font_name)

    # 查找或创建 cs
    cs = rPr.find(f"{{{A_NS}}}cs")
    if cs is None:
        cs = parse_xml(f'<a:cs {nsdecls("a")} typeface="{font_name}"/>')
        rPr.append(cs)
    else:
        cs.set("typeface", font_name)

def process_text_frame(tf, font_name="Microsoft YaHei"):
    for p in tf.paragraphs:
        p.font.name = font_name
        # 检查段落默认 rPr (defRPr)
        pPr = p._p.get_or_add_pPr()
        defRPr = pPr.find(f"{{{A_NS}}}defRPr")
        if defRPr is None:
            defRPr = parse_xml(f'<a:defRPr {nsdecls("a")}/>')
            pPr.append(defRPr)
        set_font_on_rpr(defRPr, font_name)

        for r in p.runs:
            r.font.name = font_name
            rPr = r._r.get_or_add_rPr()
            set_font_on_rpr(rPr, font_name)

def process_shape(shape, font_name="Microsoft YaHei"):
    if shape.has_text_frame:
        process_text_frame(shape.text_frame, font_name)
    elif shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                process_text_frame(cell.text_frame, font_name)
    elif shape.shape_type == 6: # Group shape
        for sub_shape in shape.shapes:
            process_shape(sub_shape, font_name)

for slide in prs.slides:
    for shape in slide.shapes:
        process_shape(shape, "Microsoft YaHei")

prs.save(v5_path)
print("Saved v5 to", v5_path)
