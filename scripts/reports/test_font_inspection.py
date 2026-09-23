import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from pptx import Presentation

v4_path = Path("/Users/guobiao/PRO/me/yoloExample/docs/presentations/2026-09-23_YOLO门店陈列识别_COO_CTO汇报_v4.pptx")
prs = Presentation(v4_path)

print(f"Total slides in modified v4: {len(prs.slides)}")

font_names = set()
for s_idx, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font and r.font.name:
                        font_names.add(r.font.name)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font and r.font.name:
                                font_names.add(r.font.name)

print("Font names found via python-pptx:", font_names)

# Check XML typeface attributes
with zipfile.ZipFile(v4_path, 'r') as z:
    for name in z.namelist():
        if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
            root = ET.fromstring(z.read(name))
            for elem in root.iter():
                if elem.tag.endswith('rPr') or elem.tag.endswith('defRPr') or elem.tag.endswith('latin') or elem.tag.endswith('ea'):
                    tf = elem.attrib.get('typeface')
                    if tf:
                        font_names.add(tf)

print("All typeface attributes across slide XMLs:", font_names)
