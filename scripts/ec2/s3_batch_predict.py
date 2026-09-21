"""在 EC2 上按 S3/Excel/JSON/TXT 清单批量下载图片并执行 YOLO 推理。

该脚本设计为由 ``scripts/ec2/s3_workflow.py`` 通过 SSH 在 EC2 项目目录中调用：

1. 先读取用户传入的图片清单，清单本身可以是本地文件、HTTP(S) URL 或 S3 URI。
2. 再把清单中的图片逐张下载到 EC2 本地工作目录。
3. 使用指定的 YOLO ``best.pt`` 做推理，保存单图 JSON、带框图片和汇总报表。
4. 最后把整个结果目录上传到用户指定的 S3 结果目录。
"""

# pylint: disable=line-too-long,too-many-branches,too-many-locals,too-many-statements,broad-exception-caught,import-error,too-many-arguments,too-many-positional-arguments

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit
from urllib.request import urlretrieve

SCRIPTS_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = SCRIPTS_ROOT.parent
if str(SCRIPTS_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_ROOT))

from common.ultralytics_config import configure_ultralytics_weights_dir  # pylint: disable=wrong-import-position

IMAGE_SUFFIXES = {".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
URL_PATTERN = re.compile(r"(?:s3|https?)://[^\s,;\"'<>]+", re.IGNORECASE)
SOURCE_FIELD_CANDIDATES = (
    "s3_uri",
    "image_url",
    "url",
    "source_url",
    "https_url",
    "public_url",
    "image",
    "path",
    "local_path",
)
PUBLIC_URL_FIELDS = ("source_url", "https_url", "public_url", "image_url", "url")


@dataclass(frozen=True)
class ImageInput:
    """单张待推理图片的清单项。

    ``source`` 是实际下载或读取的图片地址；``metadata`` 保存原始清单中的辅助字段，
    例如 JSON 清单里的 ``public_url`` 可以作为 ``s3://`` 图片的公共下载回退地址。
    """

    source: str
    row_index: int
    metadata: dict[str, Any]


@dataclass(frozen=True)
class S3Location:
    """拆解后的 S3 地址。"""

    bucket: str
    key: str


@dataclass(frozen=True)
class PredictionPaths:
    """一次批量推理的本地目录结构。"""

    root: Path
    manifest_dir: Path
    image_dir: Path
    json_dir: Path
    annotated_dir: Path


def parse_s3_uri(uri: str) -> S3Location:
    """解析 ``s3://bucket/key``，用于下载清单、图片和上传推理结果。"""
    parsed = urlsplit(uri)
    if parsed.scheme != "s3" or not parsed.netloc or not parsed.path.strip("/"):
        raise ValueError(f"S3 地址必须形如 s3://bucket/key_or_prefix：{uri}")
    return S3Location(bucket=parsed.netloc, key=parsed.path.lstrip("/"))


def boto3_client() -> Any:
    """延迟创建 boto3 S3 client，避免纯解析/单测场景强制依赖 AWS 凭证。"""
    try:
        import boto3  # pylint: disable=import-outside-toplevel
    except ImportError as exc:
        raise SystemExit("缺少 boto3，请先在 EC2 环境安装 boto3，或确认项目依赖已同步。") from exc
    return boto3.client("s3")


def ensure_prediction_paths(work_dir: Path) -> PredictionPaths:
    """创建推理工作目录，统一管理下载图、单图结果、带框图和清单缓存。"""
    root = work_dir.expanduser().resolve()
    paths = PredictionPaths(
        root=root,
        manifest_dir=root / "manifest",
        image_dir=root / "images",
        json_dir=root / "json",
        annotated_dir=root / "annotated",
    )
    for directory in (paths.root, paths.manifest_dir, paths.image_dir, paths.json_dir, paths.annotated_dir):
        directory.mkdir(parents=True, exist_ok=True)
    return paths


def download_url(source_url: str, target: Path) -> None:
    """用标准库下载 HTTP(S) 文件，避免额外依赖。"""
    try:
        urlretrieve(source_url, str(target))  # noqa: S310 - URL 来自用户清单，属于显式输入。
    except (HTTPError, URLError, OSError) as exc:
        raise RuntimeError(f"HTTP 下载失败：{source_url} -> {target}: {exc}") from exc


def download_s3_object(source_uri: str, target: Path) -> None:
    """用 boto3 下载 S3 对象到本地文件。"""
    location = parse_s3_uri(source_uri)
    boto3_client().download_file(location.bucket, location.key, str(target))


def resolve_input_manifest(source: str, manifest_dir: Path) -> Path:
    """把用户传入的清单地址解析成 EC2 本地文件路径。"""
    parsed = urlsplit(source)
    if parsed.scheme == "s3":
        suffix = Path(parsed.path).suffix or ".txt"
        target = manifest_dir / f"input_manifest{suffix}"
        download_s3_object(source, target)
        return target
    if parsed.scheme in {"http", "https"}:
        suffix = Path(parsed.path).suffix or ".txt"
        target = manifest_dir / f"input_manifest{suffix}"
        download_url(source, target)
        return target
    path = Path(source).expanduser()
    if not path.is_absolute():
        path = (PROJECT_ROOT / path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"输入清单不存在：{path}")
    return path


def extract_sources_from_text(value: object) -> list[str]:
    """从单元格或文本字段中提取图片地址。

    业务 Excel 中有时一个单元格只放一个 URL，也可能混入说明文字；因此先提取
    ``s3://`` / ``http(s)://`` 片段。若没有 URL 片段但文本看起来是本地图片路径，
    则把整个文本作为路径返回。
    """
    text = str(value or "").strip()
    if not text:
        return []
    matches = [match.group(0).strip() for match in URL_PATTERN.finditer(text)]
    if matches:
        return matches
    if Path(text).suffix.lower() in IMAGE_SUFFIXES:
        return [text]
    return []


def public_url_from_record(record: dict[str, Any]) -> str:
    """从 JSON 对象中提取可选公共 URL，供 ``s3://`` 图片在 auto/public 模式下使用。"""
    for field_name in PUBLIC_URL_FIELDS:
        value = str(record.get(field_name) or "").strip()
        if value.startswith(("http://", "https://")):
            return value
    return ""


def item_from_source(source: str, row_index: int, metadata: dict[str, Any] | None = None) -> ImageInput:
    """创建清单项，并统一裁剪空白字符。"""
    return ImageInput(source=source.strip(), row_index=row_index, metadata=metadata or {})


def sources_from_record(record: dict[str, Any], row_index: int, input_column: str) -> list[ImageInput]:
    """从 JSON 对象记录中提取图片来源。"""
    if input_column:
        if input_column not in record:
            return []
        return [item_from_source(source, row_index, dict(record)) for source in extract_sources_from_text(record[input_column])]

    metadata = dict(record)
    if record.get("s3_bucket") and record.get("s3_key"):
        source = f"s3://{record['s3_bucket']}/{str(record['s3_key']).lstrip('/')}"
        metadata.setdefault("public_url", public_url_from_record(record))
        return [item_from_source(source, row_index, metadata)]

    for field_name in SOURCE_FIELD_CANDIDATES:
        for source in extract_sources_from_text(record.get(field_name)):
            metadata.setdefault("public_url", public_url_from_record(record))
            return [item_from_source(source, row_index, metadata)]
    return []


def deduplicate_items(items: Iterable[ImageInput]) -> list[ImageInput]:
    """按图片来源去重，避免同一个清单重复推理同一张图片。"""
    seen: set[str] = set()
    unique_items: list[ImageInput] = []
    for item in items:
        if item.source in seen:
            continue
        seen.add(item.source)
        unique_items.append(item)
    return unique_items


def load_txt_manifest(path: Path) -> list[ImageInput]:
    """读取纯文本清单；每行可写一个地址，空行和注释会被忽略。"""
    items: list[ImageInput] = []
    for index, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        for source in extract_sources_from_text(stripped):
            items.append(item_from_source(source, index, {"line": index}))
    return deduplicate_items(items)


def load_csv_manifest(path: Path, input_column: str) -> list[ImageInput]:
    """读取 CSV 清单；指定列名时只读该列，否则扫描所有单元格。"""
    rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines()))
    if not rows:
        return []
    items: list[ImageInput] = []
    if input_column:
        header = [cell.strip() for cell in rows[0]]
        if input_column not in header:
            raise ValueError(f"CSV 清单中找不到列：{input_column}")
        column_index = header.index(input_column)
        for row_index, row in enumerate(rows[1:], start=2):
            value = row[column_index] if column_index < len(row) else ""
            for source in extract_sources_from_text(value):
                items.append(item_from_source(source, row_index, {"row": row_index, "column": input_column}))
        return deduplicate_items(items)

    for row_index, row in enumerate(rows, start=1):
        for value in row:
            for source in extract_sources_from_text(value):
                items.append(item_from_source(source, row_index, {"row": row_index}))
    return deduplicate_items(items)


def load_excel_manifest(path: Path, input_column: str) -> list[ImageInput]:
    """读取 xlsx/xlsm 清单；旧 xls 需先另存为 xlsx。"""
    if path.suffix.lower() == ".xls":
        raise ValueError("暂不直接读取旧版 .xls，请先另存为 .xlsx 或导出为 CSV。")
    try:
        from openpyxl import load_workbook  # pylint: disable=import-outside-toplevel
    except ImportError as exc:
        raise SystemExit("缺少 openpyxl，请先安装项目依赖。") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    items: list[ImageInput] = []
    if input_column:
        header = [str(cell or "").strip() for cell in rows[0]]
        if input_column not in header:
            raise ValueError(f"Excel 清单中找不到列：{input_column}")
        column_index = header.index(input_column)
        for row_index, row in enumerate(rows[1:], start=2):
            value = row[column_index] if column_index < len(row) else ""
            for source in extract_sources_from_text(value):
                items.append(item_from_source(source, row_index, {"row": row_index, "column": input_column}))
        return deduplicate_items(items)

    for row_index, row in enumerate(rows, start=1):
        for value in row:
            for source in extract_sources_from_text(value):
                items.append(item_from_source(source, row_index, {"row": row_index}))
    return deduplicate_items(items)


def normalize_json_records(payload: Any) -> list[Any]:
    """把多种 JSON 清单形态归一化为列表。"""
    if isinstance(payload, dict) and isinstance(payload.get("items"), list):
        return payload["items"]
    if isinstance(payload, list):
        return payload
    if isinstance(payload, dict):
        return [payload]
    raise ValueError("JSON 清单必须是数组、对象或包含 items 数组的对象。")


def load_json_manifest(path: Path, input_column: str) -> list[ImageInput]:
    """读取 JSON 清单，兼容字符串数组和对象数组。"""
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    items: list[ImageInput] = []
    for row_index, record in enumerate(normalize_json_records(payload), start=1):
        if isinstance(record, str):
            for source in extract_sources_from_text(record):
                items.append(item_from_source(source, row_index, {"index": row_index}))
        elif isinstance(record, dict):
            items.extend(sources_from_record(record, row_index, input_column))
    return deduplicate_items(items)


def load_manifest_items(path: Path, input_column: str = "") -> list[ImageInput]:
    """按扩展名分发到对应清单解析器。"""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        items = load_txt_manifest(path)
    elif suffix == ".csv":
        items = load_csv_manifest(path, input_column)
    elif suffix == ".json":
        items = load_json_manifest(path, input_column)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        items = load_excel_manifest(path, input_column)
    else:
        raise ValueError(f"不支持的清单格式：{path.suffix}，请使用 txt/csv/json/xlsx。")
    if not items:
        raise ValueError(f"清单中没有识别到图片地址：{path}")
    return items


def safe_stem(value: str, fallback: str) -> str:
    """把 URL 或文件名整理成适合落盘的短文件名。"""
    parsed = urlsplit(value)
    raw_name = Path(parsed.path or value).stem or fallback
    cleaned = re.sub(r"[^0-9A-Za-z一-鿿._-]+", "_", raw_name).strip("._-")
    return (cleaned or fallback)[:80]


def image_suffix_for_source(source: str) -> str:
    """从图片地址推断扩展名，推断失败时默认写成 jpg。"""
    parsed = urlsplit(source)
    suffix = Path(parsed.path or source).suffix.lower()
    return suffix if suffix in IMAGE_SUFFIXES else ".jpg"


def public_url_for_s3_source(source: str, item: ImageInput, public_base_url: str) -> str:
    """为 S3 图片生成可选公共 URL。"""
    explicit_url = str(item.metadata.get("public_url") or "").strip()
    if explicit_url.startswith(("http://", "https://")):
        return explicit_url
    if public_base_url:
        location = parse_s3_uri(source)
        return f"{public_base_url.rstrip('/')}/{quote(location.key, safe='/')}"
    return ""


def download_image(item: ImageInput, index: int, image_dir: Path, download_mode: str, public_base_url: str) -> Path:
    """下载或复制一张待推理图片到本地工作目录。"""
    source = item.source
    suffix = image_suffix_for_source(source)
    target = image_dir / f"{index:06d}_{safe_stem(source, f'image_{index:06d}')}{suffix}"
    if target.exists() and target.stat().st_size > 0:
        return target

    parsed = urlsplit(source)
    if parsed.scheme == "s3":
        public_url = public_url_for_s3_source(source, item, public_base_url)
        if download_mode in {"auto", "public"} and public_url:
            try:
                download_url(public_url, target)
                return target
            except RuntimeError:
                if download_mode == "public":
                    raise
        if download_mode == "public":
            raise RuntimeError(f"公共下载模式缺少可用 HTTP URL：{source}")
        download_s3_object(source, target)
        return target

    if parsed.scheme in {"http", "https"}:
        download_url(source, target)
        return target

    source_path = Path(source).expanduser()
    if not source_path.is_absolute():
        source_path = (PROJECT_ROOT / source_path).resolve()
    if not source_path.is_file():
        raise FileNotFoundError(f"图片不存在：{source_path}")
    shutil.copy2(source_path, target)
    return target


def resolve_model_path(model_path: str) -> Path:
    """解析 EC2 上的模型路径，支持项目相对路径和绝对路径。"""
    path = Path(model_path).expanduser()
    if not path.is_absolute():
        path = (PROJECT_ROOT / path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"推理模型不存在：{path}")
    return path


def prediction_json_path(json_dir: Path, image_path: Path) -> Path:
    """生成单图 JSON 输出路径。"""
    return json_dir / f"{image_path.stem}.json"


def prediction_annotated_path(annotated_dir: Path, image_path: Path) -> Path:
    """生成带框图片输出路径。"""
    return annotated_dir / f"{image_path.stem}-annotated.jpg"


def predict_one(model: Any, image_path: Path, item: ImageInput, model_path: Path, args: argparse.Namespace, paths: PredictionPaths) -> dict[str, Any]:
    """对单张图片推理并写入单图 JSON。"""
    result = model.predict(source=str(image_path), conf=args.conf, imgsz=args.imgsz, device=args.device, verbose=False)[0]
    detections = []
    class_counts: dict[str, int] = {}
    for box in result.boxes:
        class_id = int(box.cls.item())
        class_name = str(result.names[class_id])
        x1, y1, x2, y2 = (round(float(value), 2) for value in box.xyxy[0].tolist())
        class_counts[class_name] = class_counts.get(class_name, 0) + 1
        detections.append(
            {
                "class_id": class_id,
                "class_name": class_name,
                "confidence": round(float(box.conf.item()), 4),
                "xyxy": [x1, y1, x2, y2],
            }
        )

    annotated_path = prediction_annotated_path(paths.annotated_dir, image_path)
    json_path = prediction_json_path(paths.json_dir, image_path)
    result.save(filename=str(annotated_path))
    payload = {
        "source": item.source,
        "row_index": item.row_index,
        "local_image": str(image_path),
        "model": str(model_path),
        "class_counts": class_counts,
        "total_count": len(detections),
        "detections": detections,
        "annotated_image": str(annotated_path),
        "json_file": str(json_path),
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def write_summary(outputs: list[dict[str, Any]], paths: PredictionPaths, args: argparse.Namespace) -> tuple[Path, Path]:
    """写批量推理 summary.json 和 summary.csv。"""
    aggregate_counts: dict[str, int] = {}
    for output in outputs:
        for class_name, count in output.get("class_counts", {}).items():
            aggregate_counts[class_name] = aggregate_counts.get(class_name, 0) + int(count)

    success_count = sum(1 for output in outputs if output.get("status") == "success")
    failed_count = sum(1 for output in outputs if output.get("status") == "failed")
    summary = {
        "input": args.input,
        "output_s3_uri": args.output_s3_uri,
        "model": args.model,
        "conf": args.conf,
        "imgsz": args.imgsz,
        "device": args.device,
        "total_images": len(outputs),
        "success_count": success_count,
        "failed_count": failed_count,
        "aggregate_counts": aggregate_counts,
        "items": outputs,
    }
    summary_json = paths.root / "summary.json"
    summary_csv = paths.root / "summary.csv"
    summary_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    with summary_csv.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=["source", "row_index", "status", "total_count", "class_counts", "local_image", "json_file", "annotated_image", "error"],
        )
        writer.writeheader()
        for output in outputs:
            writer.writerow(
                {
                    "source": output.get("source", ""),
                    "row_index": output.get("row_index", ""),
                    "status": output.get("status", ""),
                    "total_count": output.get("total_count", ""),
                    "class_counts": json.dumps(output.get("class_counts", {}), ensure_ascii=False, sort_keys=True),
                    "local_image": output.get("local_image", ""),
                    "json_file": output.get("json_file", ""),
                    "annotated_image": output.get("annotated_image", ""),
                    "error": output.get("error", ""),
                }
            )
    return summary_json, summary_csv


def s3_key_for_output(prefix: str, root: Path, file_path: Path) -> str:
    """根据结果根目录和文件相对路径生成上传 S3 key。"""
    relative = file_path.relative_to(root).as_posix()
    return f"{prefix.rstrip('/')}/{relative}" if prefix.strip("/") else relative


def iter_result_files(paths: PredictionPaths) -> Iterable[Path]:
    """只枚举需要上传给业务查看的推理结果文件，不回传下载缓存图片。"""
    for file_path in (paths.root / "summary.json", paths.root / "summary.csv"):
        if file_path.is_file():
            yield file_path
    for directory in (paths.json_dir, paths.annotated_dir):
        yield from sorted(path for path in directory.rglob("*") if path.is_file())


def upload_results_to_s3(paths: PredictionPaths, output_s3_uri: str) -> list[str]:
    """把汇总、单图 JSON 和带框图上传到指定 S3 目录。"""
    location = parse_s3_uri(output_s3_uri)
    client = boto3_client()
    uploaded_uris = []
    for file_path in iter_result_files(paths):
        key = s3_key_for_output(location.key, paths.root, file_path)
        client.upload_file(str(file_path), location.bucket, key)
        uploaded_uris.append(f"s3://{location.bucket}/{key}")
        print(f"upload {file_path} -> s3://{location.bucket}/{key}")
    return uploaded_uris


def run_predictions(items: list[ImageInput], args: argparse.Namespace, paths: PredictionPaths) -> list[dict[str, Any]]:
    """下载图片并逐张执行 YOLO 推理；单张失败不会中断整批任务。"""
    from ultralytics import YOLO  # pylint: disable=import-outside-toplevel

    configure_ultralytics_weights_dir()
    model_path = resolve_model_path(args.model)
    model = YOLO(model_path)
    outputs: list[dict[str, Any]] = []
    for index, item in enumerate(items, start=1):
        try:
            image_path = download_image(item, index, paths.image_dir, args.download_mode, args.public_base_url.rstrip("/"))
            payload = predict_one(model, image_path, item, model_path, args, paths)
            payload["status"] = "success"
            outputs.append(payload)
            print(f"predicted {index}/{len(items)} {item.source} total_count={payload['total_count']}")
        except Exception as exc:  # pylint: disable=broad-exception-caught
            error_payload = {
                "source": item.source,
                "row_index": item.row_index,
                "status": "failed",
                "error": str(exc),
                "class_counts": {},
                "total_count": 0,
            }
            outputs.append(error_payload)
            print(f"failed {index}/{len(items)} {item.source}: {exc}")
    return outputs


def parse_args() -> argparse.Namespace:
    """解析命令行参数。"""
    parser = argparse.ArgumentParser(description="在 EC2 上按 S3/Excel/JSON/TXT 清单批量推理并上传结果。")
    parser.add_argument("--input", required=True, help="图片清单：s3://、http(s):// 或 EC2 本地文件路径")
    parser.add_argument("--output-s3-uri", required=True, help="推理结果上传目录，格式 s3://bucket/prefix")
    parser.add_argument("--model", required=True, help="EC2 上 YOLO best.pt 路径，支持项目相对路径")
    parser.add_argument("--work-dir", type=Path, default=PROJECT_ROOT / "outputs" / "ec2_predict" / "default", help="EC2 本地推理工作目录")
    parser.add_argument("--input-column", default="", help="CSV/Excel/JSON 中指定图片地址列名；留空时自动扫描")
    parser.add_argument("--download-mode", choices=["auto", "public", "boto3"], default="auto", help="S3 图片下载模式")
    parser.add_argument("--public-base-url", default="", help="公开 S3/CDN Base URL；public/auto 模式下可用")
    parser.add_argument("--conf", type=float, default=0.35, help="推理置信度阈值")
    parser.add_argument("--imgsz", type=int, default=960, help="推理图片尺寸")
    parser.add_argument("--device", default="0", help="推理设备，例如 0/cpu；空值由 Ultralytics 自动选择")
    parser.add_argument("--limit", type=int, default=0, help="只处理前 N 张图片；0 表示全量")
    return parser.parse_args()


def main() -> None:
    """脚本入口。"""
    args = parse_args()
    if not 0.0 <= args.conf <= 1.0:
        raise SystemExit("--conf 必须位于 0 到 1 之间。")
    parse_s3_uri(args.output_s3_uri)
    paths = ensure_prediction_paths(args.work_dir)
    manifest_path = resolve_input_manifest(args.input, paths.manifest_dir)
    items = load_manifest_items(manifest_path, input_column=args.input_column)
    if args.limit > 0:
        items = items[: args.limit]
    outputs = run_predictions(items, args, paths)
    summary_json, summary_csv = write_summary(outputs, paths, args)
    uploaded_uris = upload_results_to_s3(paths, args.output_s3_uri)
    result = {
        "summary_json": str(summary_json),
        "summary_csv": str(summary_csv),
        "output_s3_uri": args.output_s3_uri.rstrip("/"),
        "uploaded_files": len(uploaded_uris),
        "success_count": sum(1 for output in outputs if output.get("status") == "success"),
        "failed_count": sum(1 for output in outputs if output.get("status") == "failed"),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
